"""Structural and anonymisation checks for the Health Systems package."""
from pathlib import Path
import re
from zipfile import ZipFile

from docx import Document

ROOT = Path(__file__).resolve().parents[1]
PACKAGE = ROOT / "submission" / "Health_Systems"
REQUIRED = [
    "00_READ_ME_FIRST.txt",
    "01_FairQueue2_Manuscript_with_authors.docx",
    "02_FairQueue2_Anonymous_manuscript.docx",
    "03_FairQueue2_Title_page.docx",
    "04_FairQueue2_Cover_letter.docx",
    "05_FairQueue2_Anonymous_supplement.docx",
    "06_Health_Systems_Upload_checklist.docx",
    "figures/Figure1_workflow.tiff",
    "figures/Figure2_temporal_performance.tiff",
    "figures/Figure3_equity_utility_tradeoff.tiff",
    "tables/Table1_test_model_performance.csv",
    "tables/Table2_equity_utility_tradeoff.csv",
]


def full_text(path: Path) -> str:
    doc = Document(path)
    blocks = [p.text for p in doc.paragraphs]
    blocks.extend(cell.text for table in doc.tables for row in table.rows for cell in row.cells)
    return "\n".join(blocks)


def main() -> None:
    missing = [name for name in REQUIRED if not (PACKAGE / name).is_file()]
    assert not missing, f"Missing deliverables: {missing}"
    assert not list(PACKAGE.glob("~$*")), "Temporary Word lock file found"

    for path in PACKAGE.glob("*.docx"):
        with ZipFile(path) as archive:
            assert archive.testzip() is None, f"Corrupt DOCX member in {path.name}"

    author = full_text(PACKAGE / "01_FairQueue2_Manuscript_with_authors.docx")
    anonymous = full_text(PACKAGE / "02_FairQueue2_Anonymous_manuscript.docx")
    supplement = full_text(PACKAGE / "05_FairQueue2_Anonymous_supplement.docx")
    assert "Adebayo Aliu Adetola" in author and "Olugbenga Akinade" in author
    forbidden = ["Adebayo", "Adetola", "Akinade", "yahoo.com", "ORCID", "tollyboy88", "github.com", "zenodo"]
    for token in forbidden:
        assert token.lower() not in anonymous.lower(), f"Anonymous manuscript contains {token}"
        assert token.lower() not in supplement.lower(), f"Anonymous supplement contains {token}"

    lower = anonymous.lower()
    abstract_at = lower.index("abstract")
    practitioner_at = lower.index("practitioner summary")
    introduction_at = lower.index("1. introduction")
    assert abstract_at < practitioner_at < introduction_at
    abstract = anonymous[abstract_at:practitioner_at]
    abstract = abstract.split("Keywords:")[0]
    abstract_words = len(re.findall(r"\b[\w’'-]+\b", abstract)) - 1
    assert abstract_words <= 200, f"Abstract has {abstract_words} words"
    assert "TODO" not in author and "TBD" not in author
    print(f"PASS: {len(REQUIRED)} required files; abstract={abstract_words} words; anonymity and DOCX integrity verified")


if __name__ == "__main__":
    main()
