"""Build the submission-ready Health Systems DOCX package from verified outputs."""
from __future__ import annotations

import csv
import json
import re
from pathlib import Path
from zipfile import ZIP_DEFLATED, ZipFile

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Inches, Pt, RGBColor

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "submission" / "Health_Systems"
FIGURES = ROOT / "outputs" / "figures"
TABLES = ROOT / "outputs" / "tables"
METRICS = ROOT / "outputs" / "metrics"

TITLE = (
    "FairQueue 2.0: temporally validated forecasting and equity-constrained "
    "prioritisation of NHS elective-care pressure using public data"
)
SHORT_TITLE = "FairQueue 2.0"
NAVY = RGBColor(31, 55, 77)
TEAL = RGBColor(0, 108, 120)
GREY = RGBColor(90, 90, 90)

ABSTRACT_PARTS = [
    ("Introduction", "Long elective-care waits are unevenly distributed, while fixed weighted scores can mix prediction with normative choices and create circular targets."),
    ("Objectives", "To develop a service-level system that forecasts an independently observed outcome and exposes equity requirements as explicit policy constraints."),
    ("Methodology", "We assembled 51 monthly public NHS releases (April 2022–June 2026). Models predicted provider–specialty 52-week breach rates three months ahead. Candidate algorithms were selected chronologically and tested on 25,068 observations from nine later target months. A June 2026 top-20 illustration imposed representation floors for providers with high disparity on at least one WLMDS dimension."),
    ("Results", "Validation selected Random Forest. Test mean absolute error was 0.00684, root mean squared error 0.01434, R² 0.636 and Spearman correlation 0.836. Persistence had lower mean absolute error (0.00665), whereas Random Forest improved root mean squared error and R²; Recall@20 was identical (0.578). Raising the equity floor from the unconstrained 35% share to 40% retained Recall@20 of 0.60."),
    ("Practical implication", "Forecasting and equity choices can be audited separately, exposing who enters a monitoring list and the utility trade-off without scoring patients."),
]

PRACTITIONER = (
    "FairQueue 2.0 is designed for analysts and elective-care managers who need an auditable "
    "shortlist of provider–specialty services for further review. It forecasts the percentage "
    "of incomplete pathways expected to exceed 52 weeks three months later; it does not rank "
    "patients or recommend treatment. Managers may use the forecast-only top-20 or set a "
    "minimum representation floor for providers showing high disparity in deprivation, "
    "ethnicity, age or sex in the latest available aggregate WLMDS snapshot. The interface "
    "shows the resulting equity–utility frontier rather than hiding this choice inside a fixed "
    "weighted score. In the June 2026 retrospective illustration, increasing the high-need "
    "share from 35% to 40% left Recall@20 unchanged at 60%; this is an example, not evidence of "
    "causal or operational benefit. Before real use, organisations should confirm source-data "
    "timeliness, inspect coding changes, agree the equity definition with affected communities, "
    "and evaluate decisions prospectively alongside capacity, cost and clinical-safety evidence."
)

REFERENCES = [
    "Breiman, L. (2001) ‘Random forests’, Machine Learning, 45, pp. 5–32. https://doi.org/10.1023/A:1010933404324.",
    "British Medical Association (2025) NHS backlog data analysis. Available at: https://www.bma.org.uk/advice-and-support/nhs-delivery-and-workforce/pressures/nhs-backlog-data-analysis (accessed 5 September 2026).",
    "Chen, R.J. et al. (2023) ‘Algorithmic fairness in artificial intelligence for medicine and healthcare’, Nature Biomedical Engineering, 7, pp. 719–742. https://doi.org/10.1038/s41551-023-01056-8.",
    "Chen, T. and Guestrin, C. (2016) ‘XGBoost: a scalable tree boosting system’, Proceedings of the 22nd ACM SIGKDD International Conference on Knowledge Discovery and Data Mining, pp. 785–794. https://doi.org/10.1145/2939672.2939785.",
    "Gao, J. et al. (2025) ‘What is fair? Defining fairness in machine learning for health’, Statistics in Medicine, 44, e70234. https://doi.org/10.1002/sim.70234.",
    "Georghiou, T., Spencer, J., Scobie, S. and Raleigh, V. (2022) The elective care backlog and ethnicity. NHS Race and Health Observatory and Nuffield Trust. Available at: https://www.nhsrho.org/research/the-elective-care-backlog-and-ethnicity-2/ (accessed 5 September 2026).",
    "Gibbs, N.K. et al. (2025) ‘Prioritizing patients from the most deprived areas on elective waiting lists in the NHS in England’, MDM Policy & Practice, 10. https://doi.org/10.1177/23814683241310146.",
    "Järvelin, K. and Kekäläinen, J. (2002) ‘Cumulated gain-based evaluation of IR techniques’, ACM Transactions on Information Systems, 20, pp. 422–446. https://doi.org/10.1145/582415.582418.",
    "NHS England (2026a) Consultant-led Referral to Treatment waiting times. Available at: https://www.england.nhs.uk/statistics/statistical-work-areas/rtt-waiting-times/ (accessed 5 September 2026).",
    "NHS England (2026b) Waiting List Minimum Data Set information. Available at: https://www.england.nhs.uk/statistics/statistical-work-areas/rtt-waiting-times/wlmds/ (accessed 5 September 2026).",
    "NHS England (2026c) Monthly diagnostic waiting times and activity. Available at: https://www.england.nhs.uk/statistics/statistical-work-areas/diagnostics-waiting-times-and-activity/monthly-diagnostics-waiting-times-and-activity/ (accessed 5 September 2026).",
    "NHS England (2026d) Bed availability and occupancy. Available at: https://www.england.nhs.uk/statistics/statistical-work-areas/bed-availability-and-occupancy/ (accessed 5 September 2026).",
    "NHS England (2026e) Cancelled elective operations. Available at: https://www.england.nhs.uk/statistics/statistical-work-areas/cancelled-elective-operations/ (accessed 5 September 2026).",
    "Pedregosa, F. et al. (2011) ‘Scikit-learn: machine learning in Python’, Journal of Machine Learning Research, 12, pp. 2825–2830.",
    "Robertson, R., Blythe, N. and Jefferies, D. (2023) Tackling health inequalities on NHS waiting lists: learning from local case studies. The King’s Fund. Available at: https://www.kingsfund.org.uk/insight-and-analysis/reports/health-inequalities-nhs-waiting-lists (accessed 5 September 2026).",
    "Rudin, C. (2019) ‘Stop explaining black box machine learning models for high stakes decisions and use interpretable models instead’, Nature Machine Intelligence, 1, pp. 206–215. https://doi.org/10.1038/s42256-019-0048-x.",
]


def set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)


def add_page_number(paragraph) -> None:
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = paragraph.add_run("Page ")
    fld = OxmlElement("w:fldSimple")
    fld.set(qn("w:instr"), "PAGE")
    run._r.addnext(fld)


def configure(doc: Document, *, compact: bool = False) -> None:
    sec = doc.sections[0]
    margin = Cm(1.9 if compact else 2.4)
    sec.top_margin = margin
    sec.bottom_margin = margin
    sec.left_margin = margin
    sec.right_margin = margin
    normal = doc.styles["Normal"]
    normal.font.name = "Arial"
    normal.font.size = Pt(9.5 if compact else 10.5)
    normal._element.rPr.rFonts.set(qn("w:ascii"), "Arial")
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Arial")
    normal.paragraph_format.space_after = Pt(5)
    normal.paragraph_format.line_spacing_rule = WD_LINE_SPACING.SINGLE if compact else WD_LINE_SPACING.MULTIPLE
    if not compact:
        normal.paragraph_format.line_spacing = 1.08
    for name, size, color in (("Title", 16, NAVY), ("Heading 1", 13, NAVY), ("Heading 2", 11, TEAL), ("Heading 3", 10.5, TEAL)):
        style = doc.styles[name]
        style.font.name = "Arial"
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = color
        style._element.rPr.rFonts.set(qn("w:ascii"), "Arial")
        style._element.rPr.rFonts.set(qn("w:hAnsi"), "Arial")
        style.paragraph_format.keep_with_next = True
        style.paragraph_format.space_before = Pt(10 if name != "Title" else 0)
        style.paragraph_format.space_after = Pt(4)
    footer = sec.footer.paragraphs[0]
    add_page_number(footer)
    footer.style = doc.styles["Footer"]
    doc.core_properties.title = TITLE
    doc.core_properties.subject = "Health Systems research article submission"
    doc.core_properties.keywords = "elective care; forecasting; equity; NHS; operations research"


def add_title(doc: Document, anonymous: bool) -> None:
    p = doc.add_paragraph(style="Title")
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.add_run(TITLE)
    if not anonymous:
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = p.add_run("Adebayo Aliu Adetola¹* and Olugbenga Akinade²")
        r.bold = True
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.add_run("¹ Independent Researcher, Dewsbury, United Kingdom\n")
        p.add_run("² Teesside University, Middlesbrough, United Kingdom\n")
        p.add_run("* Corresponding author: adebayoadetola96@yahoo.com")
    else:
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.add_run("Anonymous manuscript for peer review").italic = True


def add_structured_abstract(doc: Document) -> None:
    doc.add_heading("Abstract", level=1)
    p = doc.add_paragraph()
    for label, text in ABSTRACT_PARTS:
        run = p.add_run(label + ": ")
        run.bold = True
        p.add_run(text + " ")
    p = doc.add_paragraph()
    p.add_run("Keywords: ").bold = True
    p.add_run("elective care; health equity; forecasting; operations research; waiting lists; NHS")
    doc.add_heading("Practitioner Summary", level=1)
    doc.add_paragraph(PRACTITIONER)


def add_paragraphs(doc: Document, texts: list[str]) -> None:
    for text in texts:
        doc.add_paragraph(text)


def add_table(doc: Document, headers: list[str], rows: list[list[str]], widths=None) -> None:
    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.style = "Table Grid"
    table.autofit = True
    for i, header in enumerate(headers):
        cell = table.rows[0].cells[i]
        cell.text = header
        set_cell_shading(cell, "D9EAF0")
        for run in cell.paragraphs[0].runs:
            run.bold = True
            run.font.size = Pt(8)
    for values in rows:
        cells = table.add_row().cells
        for i, value in enumerate(values):
            cells[i].text = str(value)
            cells[i].vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            for p in cells[i].paragraphs:
                for run in p.runs:
                    run.font.size = Pt(7.5)
    doc.add_paragraph()


def read_csv(name: str) -> list[dict[str, str]]:
    with (METRICS / name).open(encoding="utf-8-sig", newline="") as handle:
        return list(csv.DictReader(handle))


def add_figure(doc: Document, filename: str, caption: str) -> None:
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.add_run().add_picture(str(FIGURES / filename), width=Inches(6.15))
    c = doc.add_paragraph()
    c.alignment = WD_ALIGN_PARAGRAPH.LEFT
    r = c.add_run(caption)
    r.bold = True
    r.font.size = Pt(9)


def model_table_rows() -> list[list[str]]:
    rows = read_csv("test_model_metrics.csv")
    order = ["Persistence", "Three-month mean", "Elastic Net", "Random Forest", "Histogram Gradient Boosting", "XGBoost"]
    lookup = {row["model"]: row for row in rows}
    return [[
        model,
        f"{float(lookup[model]['mae']) * 100:.2f}",
        f"{float(lookup[model]['rmse']) * 100:.2f}",
        f"{float(lookup[model]['r2']):.3f}",
        f"{float(lookup[model]['spearman']):.3f}",
        f"{float(lookup[model]['recall_at_20']):.3f}",
        f"{float(lookup[model]['ndcg_at_20']):.3f}",
    ] for model in order]


def equity_table_rows() -> list[list[str]]:
    rows = read_csv("equity_utility_tradeoff.csv")
    picked = []
    for row in rows:
        floor = row["minimum_high_need_share"]
        if row["strategy"] == "Equity-constrained forecast" and floor not in {"0.0", "0.4", "0.5"}:
            continue
        if row["strategy"] == "Forecast risk only":
            continue
        label = row["strategy"]
        floor_label = "—" if floor == "" else f"{float(floor):.0%}"
        picked.append([
            label, floor_label, f"{float(row['high_equity_need_share']):.0%}",
            f"{float(row['observed_pressure_mean']):.2%}", f"{float(row['recall_at_20']):.2f}",
        ])
    return picked


def add_manuscript_body(doc: Document, anonymous: bool) -> None:
    doc.add_heading("1. Introduction", level=1)
    add_paragraphs(doc, [
        "England’s elective-care backlog remains a major operational and equity challenge. National totals are important, but they do not identify which provider–specialty services are most likely to experience severe future long-wait pressure. Local prioritisation is also value-laden: a list that maximises predicted pressure may under-represent services whose provider-level waiting-list composition indicates marked disparities. Public reporting therefore creates a practical operations-research problem—how to distinguish an empirical forecast from the policy choices used to act on it.",
        "An earlier FairQueue formulation combined waiting pressure, aggregate demographic disparity and operational pressure into a fixed weighted score, then trained a model to reproduce that same score. Such a model can measure fidelity to its own construction but cannot validate whether the score predicts an independently observed future outcome. Fixed weights also embed an equity judgement without showing managers the opportunity cost of alternative choices. These limitations motivated a complete redesign rather than a cosmetic revision.",
        "FairQueue 2.0 separates two questions. First, can data available at feature month t forecast the provider–specialty proportion of incomplete pathways waiting more than 52 weeks at t+3? Second, after forecast risk is estimated, how does an explicit minimum representation requirement for providers with high measured equity need change a top-K monitoring list? This architecture preserves a clear outcome, makes the normative constraint visible and retains the original weighted composite only as a comparator.",
        "The study makes four contributions. It constructs a longitudinal, versioned dataset from official aggregate NHS releases; evaluates several regression algorithms and simple temporal baselines under chronological splits; reports point, rank and provider-clustered uncertainty measures; and demonstrates an auditable equity–utility frontier. The system is deliberately service-level. It neither prioritises patients nor claims that aggregate demographic disparity causes waiting pressure.",
    ])

    doc.add_heading("2. Methods", level=1)
    doc.add_heading("2.1 Study design and unit of analysis", level=2)
    add_paragraphs(doc, [
        "We conducted a retrospective longitudinal forecasting study using monthly public aggregate data. The observational unit was one NHS provider × treatment function × feature month. The feature window ran from April 2022 to March 2026 and outcomes were observed through June 2026. We required at least 100 incomplete pathways at both feature and target months to reduce instability from small denominators.",
        "The primary outcome was the observed proportion of incomplete RTT pathways waiting more than 52 weeks exactly three calendar months after the feature month. The outcome was joined by provider code, treatment-function code and target month. This creates a target that is independent of the prioritisation score and prevents same-row label construction.",
    ])
    doc.add_heading("2.2 Data sources and provenance", level=2)
    add_paragraphs(doc, [
        "Monthly consultant-led Referral to Treatment (RTT) releases supplied incomplete pathways, counts beyond 18 and 52 weeks, decision-to-admit pathways, new RTT pathways, and admitted and non-admitted completions (NHS England, 2026a). Monthly Diagnostic Waiting Times and Activity (DM01) supplied the share waiting more than six weeks; KH03 supplied bed occupancy; and quarterly cancelled-elective-operations releases supplied cancellations and 28-day breaches (NHS England, 2026c–e). The equity stage used provider-level aggregate Waiting List Minimum Data Set (WLMDS) distributions for deprivation, ethnicity, age and sex (NHS England, 2026b). No patient-level records were used.",
        "The acquisition script records source URL, local relative path, byte size, retrieval timestamp and SHA-256 digest for each file. RTT and DM01 contributed 51 monthly files each. DM01 was lagged one month and quarterly operational variables were lagged by one completed quarter before feature joining, reflecting release availability. Missing operational values were retained for train-fitted imputation rather than filled from future observations.",
    ])
    doc.add_heading("2.3 Predictors and leakage controls", level=2)
    add_paragraphs(doc, [
        "Predictors included contemporaneous RTT rates; one-, two- and three-month lags; three- and six-month rolling summaries; backlog growth; month-on-month change; three-month trend and acceleration; demand, throughput and decision-to-admit rates; log list size; seasonality; treatment function; and lagged diagnostics, bed occupancy and cancellations. Provider identity was not used as a predictor, and WLMDS disparity variables were excluded from every forecast model.",
        "A code-level allowlist defines the feature matrix. Targets and future-dated fields are prohibited. Median imputation with missingness indicators, scaling where required and one-hot encoding were encapsulated in scikit-learn pipelines and fitted on training or development partitions only. Six automated tests check exact target shifting, split chronology, prohibited future features, operational-release lags, train-only preprocessing and the date of equity evidence.",
    ])
    doc.add_heading("2.4 Temporal partitions and candidate models", level=2)
    add_paragraphs(doc, [
        "Feature months April 2022–December 2024 formed the training partition; January–June 2025 formed validation; and July 2025–March 2026 formed the locked test partition. Their corresponding three-month-ahead target months were July 2022–March 2025, April–September 2025 and October 2025–June 2026. This is an offline chronological evaluation rather than prospective deployment. Candidate algorithm selection used validation mean absolute error (MAE), with test outcomes untouched until the model family was fixed.",
        "Baselines were persistence (the feature-month 52-week breach rate) and the mean of the previous three months. Candidate regressors were Elastic Net, Random Forest, Histogram Gradient Boosting and XGBoost. Following selection, each candidate was refitted on combined training and validation data for a common locked-test comparison. Predicted proportions were clipped to [0,1]. Hyperparameters were fixed in code rather than tuned on the test period.",
    ])
    doc.add_heading("2.5 Evaluation and uncertainty", level=2)
    add_paragraphs(doc, [
        "Point metrics were MAE, root mean squared error (RMSE) and R². Spearman correlation measured ordering. For each target month, Recall@K measured overlap between predicted and observed top-K services and normalised discounted cumulative gain (NDCG@K) rewarded correct ordering near the top (Järvelin and Kekäläinen, 2002); monthly values were averaged. We report K=10, 20 and 50 in the repository and focus on K=20 in the article.",
        "Ninety-five per cent uncertainty intervals used 100 bootstrap resamples of providers, retaining all rows for sampled providers to respect within-provider dependence. Sensitivity analyses shortened the training window to April 2023 onward, stratified results by treatment function, and compared RTT-only, RTT-plus-diagnostics and full operational feature sets.",
    ])
    doc.add_heading("2.6 Equity-constrained top-K policy", level=2)
    add_paragraphs(doc, [
        "For each of four WLMDS dimensions, a provider’s disparity statistic was converted to a national percentile. A provider was labelled high equity need when it was at or above the 75th percentile in at least one dimension. This union rule avoids assigning unverifiable weights across protected or socially relevant dimensions. In the April 2026 snapshot, 46.6% of 176 matched providers met this definition.",
        "For the June 2026 retrospective illustration, services were sorted by the validation-selected forecast and the top 20 were chosen. For minimum high-need shares q ∈ {0, 0.1, …, 0.5}, the algorithm retained the highest-risk set satisfying ceil(20q) high-need providers, replacing the lowest-risk non-high-need services only when required. We compared predicted and observed pressure, Recall@20, selected high-need share and dimension-specific disparity means. Persistence and the original 50% waiting/30% equity/20% operations composite were included as comparators, not candidate targets.",
    ])

    doc.add_heading("3. Results", level=1)
    doc.add_heading("3.1 Cohort and model selection", level=2)
    add_paragraphs(doc, [
        "The modelling cohort contained 135,708 provider–specialty–month observations from 538 providers and 23 treatment functions. The locked test contained 25,068 observations spanning nine target months. Random Forest had the lowest validation MAE among learned models and was therefore selected before test evaluation.",
        "Table 1 shows a mixed result. Random Forest achieved MAE 0.00684 (0.684 percentage points), RMSE 0.01434 (1.434 percentage points), R² 0.636 and Spearman correlation 0.836. Persistence achieved slightly lower MAE (0.00665) and nearly identical Spearman correlation, but worse RMSE (0.01575) and R² (0.561). Both attained Recall@20 of 0.578; Random Forest had higher NDCG@20 (0.784 versus 0.772). Consequently, the learned model improved large-error sensitivity and explained variance, but not every metric.",
    ])
    doc.add_paragraph("Table 1. Locked-test performance across nine target months. Errors are percentage points; the validation-selected model is Random Forest.")
    add_table(doc, ["Model", "MAE", "RMSE", "R²", "Spearman", "Recall@20", "NDCG@20"], model_table_rows())
    add_figure(doc, "Figure1_workflow.png", "Figure 1. FairQueue 2.0 analysis architecture. Future outcome forecasting and equity-constrained selection are separate stages; WLMDS indicators never enter the prediction model.")

    doc.add_heading("3.2 Temporal performance and uncertainty", level=2)
    add_paragraphs(doc, [
        "Provider-clustered bootstrap intervals for Random Forest were 0.00628–0.00740 for MAE, 0.01245–0.01651 for RMSE, 0.569–0.703 for R², 0.819–0.853 for Spearman correlation, 0.411–0.720 for Recall@20 and 0.672–0.874 for NDCG@20. Persistence intervals overlapped, supporting cautious comparison rather than a claim of unequivocal superiority.",
        "Month-specific performance varied (Figure 2), underlining the value of a temporally distributed test instead of a single snapshot. A post-April-2023 training-window sensitivity produced MAE 0.00667, RMSE 0.01415, R² 0.646 and Recall@20 0.572, close to the primary specification. Specialty-level results were heterogeneous and are supplied in the reproducibility supplement.",
    ])
    add_figure(doc, "Figure2_temporal_performance.png", "Figure 2. Monthly locked-test MAE and NDCG@20 for Random Forest and persistence, by three-month-ahead target month.")

    doc.add_heading("3.3 Ablation and equity–utility frontier", level=2)
    add_paragraphs(doc, [
        "The RTT-history-only Random Forest achieved RMSE 0.01430 and R² 0.639. Adding diagnostics yielded RMSE 0.01431 and R² 0.638; the full operational set yielded RMSE 0.01434 and R² 0.636. Operational context therefore did not materially improve aggregate point accuracy in this design, although Recall@20 increased from 0.572 to 0.578. This negative finding is retained because it constrains claims about data breadth.",
        "In the June 2026 illustration, the unconstrained forecast selected 35% high-need providers, observed mean breach rate 11.00% and Recall@20 0.60. Floors through 30% were inactive. A 40% floor changed the list while maintaining Recall@20 0.60 and observed mean breach rate 11.08%; a 50% floor also retained Recall@20 0.60 with observed mean 10.97%. Persistence obtained Recall@20 0.65 with 30% high-need representation. The legacy composite selected 75% high-need providers but Recall@20 fell to 0.15 and observed mean pressure to 3.85%, demonstrating why normative weighting should not be presented as predictive validation.",
    ])
    doc.add_paragraph("Table 2. June 2026 top-20 policy illustration. The floor is the minimum required share; achieved equity share can be higher.")
    add_table(doc, ["Policy", "Floor", "High-need share", "Observed mean", "Recall@20"], equity_table_rows())
    add_figure(doc, "Figure3_equity_utility_tradeoff.png", "Figure 3. June 2026 equity–utility frontier for the forecast-based top-20 list. Points show achieved high-equity-need representation and mean subsequently observed 52-week breach rate.")

    doc.add_heading("4. Discussion", level=1)
    doc.add_heading("4.1 Principal findings", level=2)
    add_paragraphs(doc, [
        "FairQueue 2.0 repairs the central validity problem of the earlier system by predicting a genuinely future, independently observed outcome. The learned model reduced RMSE and improved R² relative to persistence, but it did not reduce MAE and did not improve top-20 overlap. This is substantively informative: current pressure is a strong forecast of near-term pressure, and an operational tool should retain that baseline rather than imply that complexity is automatically better.",
        "The second contribution is institutional rather than purely algorithmic. Separating forecast risk from equity constraints makes the value judgement contestable. A manager can see which representation floor is active, which services enter or leave and whether observed pressure capture changes. This responds to concerns that fairness definitions in health are context-dependent (Chen et al., 2023; Gao et al., 2025) and to evidence of inequalities within elective waiting lists (Georghiou et al., 2022; Robertson et al., 2023).",
        "The June illustration showed that a modestly stronger equity floor was feasible without reducing Recall@20, but this must not be generalised beyond one target month. It is a retrospective frontier, not an intervention effect. The persistence comparator also outperformed the selected model on June Recall@20, reinforcing the need to display multiple baselines to practitioners.",
    ])
    doc.add_heading("4.2 Operational interpretation", level=2)
    add_paragraphs(doc, [
        "The intended output is a monitoring shortlist for multidisciplinary review. Forecasts could prompt investigation of capacity, demand, coding or pathway design; the system does not prescribe a response. The constraint can be adjusted only after governance bodies define what representation means locally. Any implementation should log the chosen floor, alternatives considered, resulting rank changes and subsequent actions.",
        "The ablation result suggests that more public datasets do not necessarily produce better forecasts. Release delays, provider-level aggregation and incomplete coverage may weaken operational signals. Local data with timelier theatre, staffing or referral information might add value, but should be evaluated incrementally against the same baselines and leakage controls.",
    ])
    doc.add_heading("4.3 Limitations", level=2)
    add_paragraphs(doc, [
        "The study uses aggregate pathways, not people; rates cannot support patient-level decisions or claims about individual disadvantage. Provider mergers, coding changes and treatment-function reporting can disrupt longitudinal comparability. Operational coverage was incomplete and was handled through train-fitted imputation. The WLMDS equity stage used one April 2026 provider snapshot and a deliberately simple top-quartile union rule; alternative definitions may produce different selections.",
        "Evaluation was retrospective and confined to England during a period of substantial system change. The chronological feature windows do not simulate every publication delay at a live decision date, although operational lags are applied and the outcome is always three months later. Hyperparameters were fixed rather than extensively tuned. Bootstrap resampling used 100 replicates, adequate for a reproducible uncertainty check but less precise than a larger inferential study. No prospective workflow, cost, capacity, health-outcome or equity-impact evaluation was undertaken.",
    ])

    doc.add_heading("5. Conclusion", level=1)
    doc.add_paragraph(
        "A fairness-aware prioritisation system should validate an outcome that exists independently of its policy choices. FairQueue 2.0 forecasts three-month-ahead service pressure, benchmarks the result honestly against persistence and exposes equity requirements through an explicit top-K constraint. The public-data evaluation supports the architecture and audit trail, not immediate deployment. Prospective studies should test whether the shortlist improves planning decisions and equity-relevant outcomes in practice."
    )

    doc.add_heading("Declarations", level=1)
    doc.add_heading("Data and software availability", level=2)
    if anonymous:
        doc.add_paragraph(
            "All inputs are public aggregate NHS England releases cited below. An anonymised code-and-output archive is supplied with the submission; the public repository and archival identifiers are withheld during peer review and will be restored in the accepted version. The source manifest records URLs, retrieval times and SHA-256 digests."
        )
    else:
        doc.add_paragraph(
            "All inputs are public aggregate NHS England releases cited below. Reproducible code, tests, source manifest and derived non-sensitive outputs are available at https://github.com/tollyboy88/fairqueue-simulator. The earlier software archive is https://doi.org/10.5281/zenodo.21347299 and processed-results archive is https://doi.org/10.5281/zenodo.21347317; a versioned FairQueue 2.0 release should be minted upon acceptance."
        )
    doc.add_heading("Ethics statement", level=2)
    doc.add_paragraph("The study analysed public, aggregate service statistics and did not involve human participants, identifiable records or an intervention; research ethics approval and consent were therefore not applicable.")
    doc.add_heading("Author contributions", level=2)
    if anonymous:
        doc.add_paragraph("Author contribution details are provided on the separate title page and omitted here for anonymous review.")
    else:
        doc.add_paragraph("A.A.A.: Conceptualisation, data curation, formal analysis, investigation, methodology, software, validation, visualisation, writing—original draft. O.A.: Supervision, methodology, validation, writing—review and editing. Both authors approved the submitted version.")
    doc.add_heading("Funding", level=2)
    doc.add_paragraph("The authors received no specific grant for this research.")
    doc.add_heading("Competing interests", level=2)
    doc.add_paragraph("The authors declare no competing interests.")
    doc.add_heading("Acknowledgements", level=2)
    if anonymous:
        doc.add_paragraph("Acknowledgements are omitted for anonymous review.")
    else:
        doc.add_paragraph("The authors thank NHS England for making the aggregate source data openly available.")

    doc.add_heading("References", level=1)
    for ref in REFERENCES:
        p = doc.add_paragraph(ref)
        p.paragraph_format.left_indent = Cm(0.6)
        p.paragraph_format.first_line_indent = Cm(-0.6)
        p.paragraph_format.space_after = Pt(3)


def build_manuscript(path: Path, anonymous: bool) -> None:
    doc = Document()
    configure(doc)
    doc.core_properties.author = "" if anonymous else "Adebayo Aliu Adetola; Olugbenga Akinade"
    doc.core_properties.last_modified_by = "" if anonymous else "FairQueue authors"
    add_title(doc, anonymous)
    add_structured_abstract(doc)
    add_manuscript_body(doc, anonymous)
    doc.save(path)


def build_title_page(path: Path) -> None:
    doc = Document()
    configure(doc)
    doc.core_properties.author = "Adebayo Aliu Adetola; Olugbenga Akinade"
    add_title(doc, False)
    doc.add_heading("Article information", level=1)
    add_table(doc, ["Item", "Information"], [
        ["Article type", "Research Article"],
        ["Target journal", "Health Systems (Taylor & Francis / OR Society)"],
        ["Running title", SHORT_TITLE],
        ["Keywords", "elective care; health equity; forecasting; operations research; waiting lists; NHS"],
        ["Figures and tables", "3 figures and 2 tables in the manuscript"],
        ["Corresponding author", "Adebayo Aliu Adetola, adebayoadetola96@yahoo.com, Dewsbury, United Kingdom"],
    ])
    doc.add_heading("Author identifiers", level=1)
    doc.add_paragraph("Adebayo Aliu Adetola — ORCID 0009-0002-1995-6400")
    doc.add_paragraph("Olugbenga Akinade — ORCID 0000-0003-3950-3775")
    doc.add_heading("Contributions", level=1)
    doc.add_paragraph("A.A.A.: Conceptualisation, data curation, formal analysis, investigation, methodology, software, validation, visualisation, writing—original draft. O.A.: Supervision, methodology, validation, writing—review and editing. Both authors approved the submitted version.")
    doc.add_heading("Declarations", level=1)
    doc.add_paragraph("Funding: The authors received no specific grant for this research.")
    doc.add_paragraph("Competing interests: The authors declare no competing interests.")
    doc.add_paragraph("Ethics: Public aggregate data only; ethics approval and consent were not applicable.")
    doc.add_paragraph("Data/software: https://github.com/tollyboy88/fairqueue-simulator. Versioned archives should be updated to FairQueue 2.0 upon acceptance.")
    doc.save(path)


def build_cover_letter(path: Path) -> None:
    doc = Document()
    configure(doc)
    doc.core_properties.author = "Adebayo Aliu Adetola"
    p = doc.add_paragraph()
    p.add_run("5 September 2026\n")
    p.add_run("Editors-in-Chief\nHealth Systems")
    doc.add_heading("Submission of a Research Article", level=1)
    doc.add_paragraph("Dear Editors-in-Chief,")
    doc.add_paragraph(f"Please consider our manuscript, “{TITLE}”, for publication as a Research Article in Health Systems.")
    add_paragraphs(doc, [
        "The paper addresses a health-service operations problem at the intersection of forecasting, transparent decision support and health equity. It uses 51 monthly public NHS releases to predict provider–specialty 52-week breach rates three months ahead, evaluates candidate models on a chronological holdout, and then applies equity requirements as explicit top-K representation constraints. This separation avoids treating a researcher-authored priority score as its own validation target.",
        "The results are deliberately balanced. Random Forest improved test RMSE and R² relative to persistence, but persistence retained slightly lower MAE and identical Recall@20. The policy illustration shows how a 40% high-equity-need floor changed the June 2026 top-20 list without reducing observed Recall@20, while making clear that this is retrospective and not causal evidence. We believe this combination of operations research, data science and implementable governance is directly aligned with Health Systems’ interest in system-level methods that improve healthcare delivery.",
        "The manuscript is original, is not under consideration elsewhere, and has been approved by both authors. The study used public aggregate data only. The authors declare no competing interests and no specific funding. An anonymised manuscript and supplementary reproducibility file are included for review, alongside a separate title page and unblinded copy.",
        "Thank you for considering the work. Correspondence should be addressed to Adebayo Aliu Adetola at adebayoadetola96@yahoo.com.",
    ])
    doc.add_paragraph("Yours sincerely,")
    doc.add_paragraph("Adebayo Aliu Adetola\nOn behalf of both authors")
    doc.save(path)


def build_supplement(path: Path) -> None:
    doc = Document()
    configure(doc, compact=True)
    doc.core_properties.author = ""
    doc.core_properties.last_modified_by = ""
    doc.add_heading("FairQueue 2.0 — Supplementary methods and reproducibility", level=0)
    doc.add_paragraph("Anonymous supplementary file for peer review. This supplement contains no author names or repository identity.")
    doc.add_heading("S1. Data lineage and temporal contract", level=1)
    doc.add_paragraph("The acquisition manifest contains 102 RTT and DM01 monthly source files (51 per collection) plus referenced quarterly operational and WLMDS sources. Every downloaded file is paired with its official URL, byte count, UTC retrieval time and SHA-256 digest. The analysis uses feature months April 2022–March 2026 and outcomes exactly three calendar months later through June 2026.")
    add_table(doc, ["Partition", "Feature months", "Target months", "Role"], [
        ["Training", "Apr 2022–Dec 2024", "Jul 2022–Mar 2025", "Fit preprocessors and candidates"],
        ["Validation", "Jan–Jun 2025", "Apr–Sep 2025", "Select model by MAE"],
        ["Locked test", "Jul 2025–Mar 2026", "Oct 2025–Jun 2026", "Final evaluation only"],
    ])
    doc.add_heading("S2. Model specifications", level=1)
    add_table(doc, ["Model", "Fixed specification"], [
        ["Elastic Net", "α=0.001; l1 ratio=0.25; max_iter=20,000; scaled numeric features"],
        ["Random Forest", "60 trees; depth=12; min leaf=20; max features=0.6; seed=42"],
        ["Histogram GB", "250 iterations; learning rate=0.05; 31 leaves; min leaf=25; L2=1"],
        ["XGBoost", "300 trees; depth=5; learning rate=0.05; row/column sample=0.85; seed=42"],
        ["Preprocessing", "Training-fitted median imputation + missing indicators; one-hot specialty"],
    ])
    doc.add_heading("S3. Integrity checks", level=1)
    doc.add_paragraph("The supplied automated suite passed 6/6 checks: exact t+3 target join; chronological feature partitions; no target/future columns in the allowlist; one-quarter lag for quarterly sources; train-only preprocessing; and equity snapshot date preceding the evaluated June 2026 outcome.")
    doc.add_heading("S4. Uncertainty and sensitivity", level=1)
    add_table(doc, ["Analysis", "MAE", "RMSE", "R²", "Recall@20"], [
        ["Primary Random Forest", "0.00684", "0.01434", "0.636", "0.578"],
        ["Provider bootstrap 95% CI", "0.00628–0.00740", "0.01245–0.01651", "0.569–0.703", "0.411–0.720"],
        ["Training from Apr 2023", "0.00667", "0.01415", "0.646", "0.572"],
        ["RTT-history features only", "0.00683", "0.01430", "0.639", "0.572"],
        ["RTT + diagnostics", "0.00684", "0.01431", "0.638", "0.578"],
        ["Full operational context", "0.00684", "0.01434", "0.636", "0.578"],
    ])
    doc.add_heading("S5. Equity policy algorithm", level=1)
    doc.add_paragraph("High equity need is a provider-level Boolean: at or above the national 75th percentile for any one of deprivation, ethnicity, age or sex disparity. For top-K size 20 and requested floor q, the algorithm requires ceil(20q) high-need providers. If the unconstrained list misses the requirement, it replaces the lowest-forecast non-high-need services with the highest-forecast eligible services outside the list until the floor is satisfied. Dimension values and every selected row are exported for audit.")
    doc.add_heading("S6. Reproduction sequence", level=1)
    doc.add_paragraph("Install requirements, run `py run_all.py` (or `--skip-download` when manifest files already exist), then run `py -m pytest -q`. Generated metrics, selections, three figures and two tables are deterministic under seed 42. The anonymous code archive supplied to reviewers mirrors the public repository but omits identity-bearing metadata.")
    doc.save(path)


def build_checklist(path: Path) -> None:
    doc = Document()
    configure(doc)
    doc.core_properties.author = "Adebayo Aliu Adetola"
    doc.add_heading("Health Systems submission upload checklist", level=0)
    doc.add_paragraph("Prepared 5 September 2026. Verify the live ScholarOne fields immediately before clicking Submit, as portal labels can change.")
    doc.add_heading("Recommended upload order", level=1)
    add_table(doc, ["Order", "Portal designation", "File", "Visible to reviewers?"], [
        ["1", "Main Document / Manuscript", "02_FairQueue2_Anonymous_manuscript.docx", "Yes"],
        ["2", "Title Page", "03_FairQueue2_Title_page.docx", "No"],
        ["3", "Cover Letter", "04_FairQueue2_Cover_letter.docx", "No"],
        ["4", "Supplemental Material for Review", "05_FairQueue2_Anonymous_supplement.docx", "Yes"],
        ["5", "Figure", "figures/Figure1_workflow.tiff", "Yes"],
        ["6", "Figure", "figures/Figure2_temporal_performance.tiff", "Yes"],
        ["7", "Figure", "figures/Figure3_equity_utility_tradeoff.tiff", "Yes"],
        ["8", "Manuscript—unblinded / not for review", "01_FairQueue2_Manuscript_with_authors.docx", "No"],
    ])
    doc.add_paragraph("The two CSV files in `tables/` are audit-friendly copies; the same two tables are already embedded in both manuscripts. Upload them only if the portal requests separate table files.")
    doc.add_heading("Verified article requirements", level=1)
    for item in [
        "Article type: Research Article; structured abstract uses Introduction, Objectives, Methodology, Results and Practical implication.",
        "Abstract is under the journal’s 200-word limit; six keywords are supplied (allowed range: 3–7).",
        "The article is below the approximate 8,000-word guidance and contains three figures plus two tables.",
        "Practitioner Summary appears between the abstract material and Introduction in both manuscript versions.",
        "Anonymous files contain no author names, affiliations, email addresses, ORCIDs or public-repository identity; personal metadata has been scrubbed.",
        "Subscription publication can be selected without an article publishing charge; open access is optional and should be chosen only if funding is available.",
    ]:
        doc.add_paragraph(item, style="List Bullet")
    doc.add_heading("ScholarOne form entries", level=1)
    doc.add_paragraph("Copy the title and structured abstract from the manuscript. Enter all authors and ORCIDs from the title page. Select 3–7 matching keywords. Declare: no specific funding; no competing interests; public aggregate data only; no ethics approval required. Suggest reviewers only if the portal requests them and after checking institutional conflicts.")
    doc.add_heading("Final pre-submit checks", level=1)
    for item in [
        "Open the generated PDF proof and confirm figure/table placement.",
        "Confirm the anonymous manuscript—not the author copy—is designated for review.",
        "Confirm no tracked changes, comments, hidden author details or temporary Word lock files remain.",
        "Confirm the GitHub repository and any Zenodo archive show the same FairQueue 2.0 code/results before making them public in the unblinded version.",
        "Keep the ZIP as an archive; upload individual files because ScholarOne assigns a designation to each file.",
    ]:
        doc.add_paragraph(item, style="List Bullet")
    doc.add_heading("Official guidance checked", level=1)
    doc.add_paragraph("Journal instructions: https://www.tandfonline.com/journals/thss20/about-this-journal")
    doc.add_paragraph("Taylor & Francis submission checklist: https://authorservices.taylorandfrancis.com/publishing-your-research/making-your-submission/article-submission-checklist/")
    doc.add_paragraph("Anonymous peer review guidance: https://authorservices.taylorandfrancis.com/publishing-your-research/peer-review/anonymous-peer-review/")
    doc.save(path)


def count_words(path: Path) -> int:
    doc = Document(path)
    text = " ".join(p.text for p in doc.paragraphs)
    for table in doc.tables:
        text += " " + " ".join(cell.text for row in table.rows for cell in row.cells)
    return len(re.findall(r"\b[\w’'-]+\b", text))


def main() -> None:
    OUT.mkdir(parents=True, exist_ok=True)
    (OUT / "figures").mkdir(exist_ok=True)
    (OUT / "tables").mkdir(exist_ok=True)
    paths = {
        "01_FairQueue2_Manuscript_with_authors.docx": lambda p: build_manuscript(p, False),
        "02_FairQueue2_Anonymous_manuscript.docx": lambda p: build_manuscript(p, True),
        "03_FairQueue2_Title_page.docx": build_title_page,
        "04_FairQueue2_Cover_letter.docx": build_cover_letter,
        "05_FairQueue2_Anonymous_supplement.docx": build_supplement,
        "06_Health_Systems_Upload_checklist.docx": build_checklist,
    }
    for filename, builder in paths.items():
        builder(OUT / filename)
    for source in FIGURES.glob("*.tiff"):
        (OUT / "figures" / source.name).write_bytes(source.read_bytes())
    for source in (TABLES / "Table1_test_model_performance.csv", TABLES / "Table2_equity_utility_tradeoff.csv"):
        (OUT / "tables" / source.name).write_bytes(source.read_bytes())
    notes = [
        "HEALTH SYSTEMS SUBMISSION PACKAGE",
        "Upload individual files in the numbered order described in 06_Health_Systems_Upload_checklist.docx.",
        "Use 02_FairQueue2_Anonymous_manuscript.docx as the review manuscript.",
        "Do not designate 01_FairQueue2_Manuscript_with_authors.docx or the title page for reviewer access.",
        "Tables are embedded; separate CSVs are optional unless ScholarOne requests them.",
        "The ZIP is an archive, not the preferred ScholarOne upload object.",
    ]
    (OUT / "00_READ_ME_FIRST.txt").write_text("\n\n".join(notes) + "\n", encoding="utf-8")
    archive = OUT / "FairQueue2_Health_Systems_Submission.zip"
    with ZipFile(archive, "w", ZIP_DEFLATED) as zf:
        for item in sorted(OUT.rglob("*")):
            if item.is_file() and item != archive and "_rendered" not in item.parts:
                zf.write(item, item.relative_to(OUT))
    print("Structured abstract words:", sum(len(re.findall(r"\b[\w’'-]+\b", t)) for _, t in ABSTRACT_PARTS))
    for filename in paths:
        print(filename, count_words(OUT / filename), "words")
    print("Wrote", archive)


if __name__ == "__main__":
    main()
